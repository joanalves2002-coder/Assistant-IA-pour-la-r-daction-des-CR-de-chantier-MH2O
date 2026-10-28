import streamlit as st
import anthropic
import re
from docx import Document
from docx.shared import Pt, RGBColor
import io
from datetime import date

# ── Config page ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Assistant CR — MH2O",
    page_icon="🏗️",
    layout="wide"
)

# ── Style ─────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
    .main { background-color: #f4f7fb; }
    .stApp { font-family: 'DM Sans', sans-serif; }
    .header-box {
        background: #0d1f3c;
        color: white;
        padding: 24px 32px;
        border-radius: 12px;
        margin-bottom: 24px;
    }
    .header-box h1 { color: white; margin: 0; font-size: 28px; }
    .header-box p { color: rgba(255,255,255,0.7); margin: 8px 0 0; }
    .accent { color: #e8a020; }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="header-box">
    <h1>🏗️ Assistant CR <span class="accent">MH2O</span></h1>
    <p>Reformulation de notes de chantier et insertion automatique dans le Word du CR précédent</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("### ⚙️ Configuration")
    api_key = st.text_input("Clé API Claude", type="password", placeholder="sk-ant-...")
    st.markdown("---")
    st.markdown("""
### 📖 Comment utiliser
1. Entre ta clé API Claude
2. Upload le Word du CR précédent
3. Renseigne les infos du nouveau CR
4. Choisis la section à mettre à jour
5. Saisis tes notes brutes
6. Génère et télécharge le nouveau CR
    """)

# ── Constantes ────────────────────────────────────────────────────────────────
BLEU = RGBColor(0x00, 0x00, 0xFF)
BLEU_TITRE = RGBColor(0x98, 0xB5, 0xD8)
NOIR = RGBColor(0x00, 0x00, 0x00)

# ── Fonctions ─────────────────────────────────────────────────────────────────

def ecrire_cellule_texte(cell, texte, nouveau=True):
    """Calibri 8pt, bleu si nouveau, entreprises (MAJUSCULES) en gras."""
    para = cell.paragraphs[0]
    # Vider les runs existants
    for run in para.runs:
        run.text = ""
    # Découper : mots tout en majuscules vs reste
    tokens = re.split(r'(\b[A-Z]{2}[A-Z0-9\-&/]*\b)', texte)
    couleur = BLEU if nouveau else NOIR
    for token in tokens:
        if not token:
            continue
        run = para.add_run(token)
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = couleur
        run.bold = bool(re.match(r'^[A-Z]{2}[A-Z0-9\-&/]*$', token))

def ecrire_cellule_simple(cell, texte, nouveau=True):
    """Calibri 8pt, bleu si nouveau."""
    para = cell.paragraphs[0]
    for run in para.runs:
        run.text = ""
    run = para.add_run(texte)
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = BLEU if nouveau else NOIR

def parse_ligne(ligne):
    """Extrait (texte_propre, statut, responsable) d'une ligne de CR."""
    ligne = ligne.strip()
    ligne = re.sub(r'^[-•*]\s*', '', ligne)

    responsable = ""
    statut = ""

    # Responsable : dernier (contenu) en fin de ligne
    m = re.search(r'\(([^)]+)\)\s*$', ligne)
    if m:
        responsable = m.group(1).strip()
        ligne = ligne[:m.start()].strip().rstrip('.')

    # Statut : PA, PM, EN ATT, FAIT, ou date DD/MM/YY
    for pat in [r'\b(PA|PM|EN ATT|EN ATTENTE|FAIT)\b', r'\b(\d{2}/\d{2}/\d{2,4})\b']:
        ms = re.search(pat, ligne, re.IGNORECASE)
        if ms:
            statut = ms.group(1).upper()
            ligne = (ligne[:ms.start()] + ligne[ms.end():]).strip().rstrip('.')
            break

    return ligne.strip(), statut, responsable


# ── Upload ────────────────────────────────────────────────────────────────────
st.markdown("### 📂 Étape 1 — Upload du CR précédent")
uploaded_file = st.file_uploader(
    "Dépose le fichier Word du CR précédent (.docx)",
    type=["docx"]
)

if uploaded_file:
    doc_bytes = uploaded_file.read()
    doc = Document(io.BytesIO(doc_bytes))

    # Détecter les sections (tables ≥5 colonnes, hors DESIGNATION)
    sections = {}
    for i, table in enumerate(doc.tables):
        if len(table.columns) >= 5 and len(table.rows) >= 1:
            en_tete = table.rows[0].cells[0].text.strip()
            if en_tete and len(en_tete) > 2 and "DESIGNATION" not in en_tete.upper():
                label = f"{en_tete[:60]} (tableau {i+1})"
                sections[label] = i

    st.success(f"✅ {len(doc.tables)} tableaux trouvés · {len(sections)} sections détectées")

    # ── Infos CR ─────────────────────────────────────────────────────────────
    st.markdown("### 📋 Étape 2 — Informations du nouveau CR")
    c1, c2, c3 = st.columns(3)
    with c1:
        num_cr = st.text_input("Numéro du CR", placeholder="Ex : 29")
    with c2:
        date_cr = st.date_input("Date de la réunion", value=date.today())
    with c3:
        prochain = st.text_input("Prochain RDV", placeholder="Ex : MARDI 03 NOVEMBRE 2026 à 09h00")

    # ── Section ───────────────────────────────────────────────────────────────
    st.markdown("### 🎯 Étape 3 — Section à mettre à jour")
    if sections:
        section = st.selectbox("Section :", options=list(sections.keys()))
    else:
        st.warning("⚠️ Aucune section détectée.")
        section = None

    # ── Notes ─────────────────────────────────────────────────────────────────
    st.markdown("### ✏️ Étape 4 — Notes brutes")
    notes = st.text_area(
        "Tes notes de réunion :",
        height=200,
        placeholder="- serbois retard bat3\n- bazzi peinture 80% bat1\n- enedis raccordement prévu 3 avril"
    )

    # ── Génération ────────────────────────────────────────────────────────────
    if st.button("🤖 Générer les points CR avec l'IA", type="primary", use_container_width=True):
        if not api_key:
            st.error("⚠️ Entre ta clé API dans la barre latérale")
        elif not notes.strip():
            st.error("⚠️ Saisis tes notes")
        else:
            with st.spinner("L'IA reformule tes notes..."):
                try:
                    client = anthropic.Anthropic(api_key=api_key)
                    prompt = f"""Tu es assistant MOE chez MH2O, ingénierie du bâtiment.
Reformule ces notes brutes en points de CR professionnels.

RÈGLES :
- Chaque point = une phrase courte et factuelle
- Noms d'entreprises en MAJUSCULES (SERBOIS, BAZZI, ENEDIS, ROC...)
- Statut en fin : PA (Pour Action), PM (Pour Mémoire), EN ATT (En Attente), ou date DD/MM/YY si échéance
- Responsable entre parenthèses : PA (SERBOIS)
- Format : Texte. STATUT (RESPONSABLE)
- Pas d'introduction, commence directement par les points

EXEMPLES :
MH2O demande à SERBOIS de transmettre un planning actualisé pour le Bâtiment 3. PA (SERBOIS)
Raccordement ENEDIS prévu le 03/04/26, confirmation attendue. EN ATT (MH2O)
Avancement peinture BAZZI : 80% Bâtiment 1, retard de 2 semaines. PA (BAZZI)

Notes :
{notes}

Section : {section}
CR n° : {num_cr} — Date : {date_cr.strftime('%d/%m/%Y')}"""

                    msg = client.messages.create(
                        model="claude-sonnet-4-6",
                        max_tokens=1500,
                        messages=[{"role": "user", "content": prompt}]
                    )
                    st.session_state.update({
                        "points": msg.content[0].text,
                        "section": section,
                        "sections": sections,
                        "num_cr": num_cr,
                        "date_cr": date_cr,
                        "prochain": prochain,
                        "doc_bytes": doc_bytes,
                    })
                except Exception as e:
                    st.error(f"Erreur API : {e}")

    # ── Résultat & export ─────────────────────────────────────────────────────
    if "points" in st.session_state:
        st.markdown("### ✅ Étape 5 — Points générés")
        points_edites = st.text_area(
            "Vérifie et corrige si besoin :",
            value=st.session_state["points"],
            height=280,
            key="points_edites"
        )

        st.markdown("### 📥 Étape 6 — Générer le Word")
        if st.button("📄 Créer le nouveau CR Word", type="primary", use_container_width=True):
            with st.spinner("Insertion dans le document..."):
                try:
                    doc2 = Document(io.BytesIO(st.session_state["doc_bytes"]))
                    num = st.session_state["num_cr"]
                    dt = st.session_state["date_cr"]
                    rdv = st.session_state["prochain"]
                    smap = st.session_state["sections"]
                    sel = st.session_state["section"]
                    date_fr = dt.strftime("%d/%m/%Y")

                    # Jours en français pour le titre
                    jours_fr = ["LUNDI", "MARDI", "MERCREDI", "JEUDI", "VENDREDI", "SAMEDI", "DIMANCHE"]
                    mois_fr = ["JANVIER", "FÉVRIER", "MARS", "AVRIL", "MAI", "JUIN",
                               "JUILLET", "AOÛT", "SEPTEMBRE", "OCTOBRE", "NOVEMBRE", "DÉCEMBRE"]
                    jour_nom = jours_fr[dt.weekday()]
                    mois_nom = mois_fr[dt.month - 1]
                    date_titre = f"{jour_nom} {dt.day} {mois_nom} {dt.year}"

                    # Collecter TOUS les paragraphes : corps + cellules de tableaux
                    tous_paragraphes = list(doc2.paragraphs)
                    for tbl in doc2.tables:
                        for row in tbl.rows:
                            for cell in row.cells:
                                for cp in cell.paragraphs:
                                    tous_paragraphes.append(cp)

                    for para in tous_paragraphes:
                        txt = para.text.upper()

                        # Mise à jour numéro + date dans le titre principal
                        if "CR DE REUNION" in txt or "CR DE RÉUNION" in txt:
                            for run in para.runs:
                                # Remplacer le numéro (N°28 → N°{num})
                                if num:
                                    run.text = re.sub(r'N°\s*\d+', f'N°{num}', run.text)
                                # Remplacer "DU MARDI 20 OCTOBRE 2026"
                                run.text = re.sub(
                                    r'\bDU\s+(LUNDI|MARDI|MERCREDI|JEUDI|VENDREDI|SAMEDI|DIMANCHE)\s+\d+\s+\w+\s+\d{4}',
                                    f'DU {date_titre}', run.text, flags=re.IGNORECASE
                                )
                                # Remplacer "MARDI 20 OCTOBRE 2026" sans DU
                                run.text = re.sub(
                                    r'(LUNDI|MARDI|MERCREDI|JEUDI|VENDREDI|SAMEDI|DIMANCHE)\s+\d+\s+\w+\s+\d{4}',
                                    date_titre, run.text, flags=re.IGNORECASE
                                )

                        # Mise à jour prochain RDV
                        if "PROCHAIN RDV" in txt and rdv:
                            full_text = para.text
                            # Remplacer tout ce qui suit "le "
                            new_text = re.sub(
                                r'(?<=\ble\s).*',
                                rdv, full_text, flags=re.IGNORECASE | re.DOTALL
                            )
                            if new_text == full_text:
                                # Fallback : depuis le premier jour de semaine
                                new_text = re.sub(
                                    r'(LUNDI|MARDI|MERCREDI|JEUDI|VENDREDI|SAMEDI|DIMANCHE).*',
                                    rdv, full_text, flags=re.IGNORECASE | re.DOTALL
                                )
                            if new_text != full_text and para.runs:
                                para.runs[0].text = new_text
                                for run in para.runs[1:]:
                                    run.text = ""

                    # Insertion dans le tableau
                    if sel and sel in smap:
                        table = doc2.tables[smap[sel]]
                        nb_cols = len(table.columns)

                        for ligne in points_edites.strip().split('\n'):
                            ligne = ligne.strip()
                            if not ligne or ligne.startswith('#'):
                                continue
                            texte, statut, resp = parse_ligne(ligne)
                            row = table.add_row()
                            cells = row.cells

                            ecrire_cellule_texte(cells[0], texte, nouveau=True)
                            if nb_cols > 1:
                                ecrire_cellule_simple(cells[1], resp, nouveau=True)
                            if nb_cols > 2:
                                ecrire_cellule_simple(cells[2], date_fr, nouveau=True)
                            if nb_cols > 3:
                                ecrire_cellule_simple(cells[3], statut, nouveau=True)
                            if nb_cols > 4:
                                ecrire_cellule_simple(cells[4], "", nouveau=False)

                    output = io.BytesIO()
                    doc2.save(output)
                    output.seek(0)
                    nom = f"CR_CHANTIER_N{num}_{dt.strftime('%d-%m-%Y')}.docx"
                    st.download_button(
                        f"⬇️ Télécharger {nom}",
                        data=output,
                        file_name=nom,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
                    st.success("✅ CR généré ! Télécharge le fichier ci-dessus.")

                except Exception as e:
                    st.error(f"Erreur : {e}")
                    import traceback
                    st.code(traceback.format_exc())

else:
    st.info("👆 Upload le fichier Word du CR précédent pour commencer")
